"""Write the final consolidated Step-3 supplementary workbook (Tables S35-S42)."""
import pickle
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from biabak import config as C

OUT = C.OUT_DIR / "step3"

A = pickle.load(open(OUT / "step3a_tables.pkl", "rb"))
B = pickle.load(open(OUT / "step3b_tables.pkl", "rb"))
C3 = pickle.load(open(OUT / "step3c_tables.pkl", "rb"))

# ---------------------------------------------------------------------
# Final benchmark design
# ---------------------------------------------------------------------
design = B["benchmark_design"].copy()

mask = design["element"].astype(str).str.lower().eq("replicates per scenario")
design.loc[mask, "setting"] = (
    "initial benchmark: 30 replicates per scenario; "
    "final covariate-only benchmark after Step 3C: "
    "330 replicates for a = 0 and 130 replicates for each non-zero "
    "covariate share (0.05, 0.10, 0.20, 0.35); "
    "range-sensitivity extension: 40 replicates per tested combination "
    "at 5 and 40 km"
)

# ---------------------------------------------------------------------
# Consolidated replicate-level results
# ---------------------------------------------------------------------
initial_reps = B["benchmark_all_replicates"].copy()
initial_reps["source"] = "initial_step3b"

extension_reps = C3["extension_replicates"].copy()
extension_reps["source"] = "extension_step3c"

all_replicates = pd.concat(
    [initial_reps, extension_reps],
    ignore_index=True,
    sort=False
)

# ---------------------------------------------------------------------
# Final tables S35-S42
# ---------------------------------------------------------------------
order = [
    (
        A["aoa_summary"].T.reset_index().rename(
            columns={"index": "quantity", 0: "value"}
        ),
        "Area of applicability: feature space, threshold and coverage of the Centre Region",
    ),
    (
        A["aoa_by_distance"],
        "Share of the prediction grid inside the area of applicability by distance to the nearest borehole",
    ),
    (
        A["aoa_training_DI"],
        "Training dissimilarity index of every borehole under its NNDM fold",
    ),
    (
        design,
        "Design of the synthetic detection benchmark, including the final Step 3C replication extension",
    ),
    (
        C3["power"],
        "Final covariate-signal detection rates, Wilson intervals, observed statistics and one-sided compatibility results",
    ),
    (
        C3["spatial"],
        "Final spatial-signal detection rates by spatial share and range parameter",
    ),
    (
        B["real_data_same_evaluators"],
        "Observed archive analysed with the same learners and tests used in the synthetic benchmark",
    ),
    (
        all_replicates,
        "Every replicate of the initial and extended synthetic benchmark",
    ),
]

path = OUT / "Supplementary_Tables_Step3_Applicability_Benchmark.xlsx"

readme = []

with pd.ExcelWriter(path, engine="openpyxl") as xw:
    for n, (df, title) in enumerate(order, start=35):
        df = df.copy()

        for c in df.columns:
            if df[c].dtype == bool:
                df[c] = df[c].map({True: "yes", False: "no"})
            if str(df[c].dtype) == "category":
                df[c] = df[c].astype(str)

        sheet = f"S{n}"

        pd.DataFrame(
            [[f"Table S{n}. {title}"]]
        ).to_excel(
            xw,
            sheet_name=sheet,
            index=False,
            header=False,
        )

        df.to_excel(
            xw,
            sheet_name=sheet,
            index=False,
            startrow=2,
        )

        readme.append(
            dict(
                table=f"Table S{n}",
                sheet=sheet,
                title=title,
                rows=len(df),
                columns=len(df.columns),
            )
        )

    pd.DataFrame(readme).to_excel(
        xw,
        sheet_name="README",
        index=False,
    )

wb = load_workbook(path)
wb.move_sheet("README", offset=-len(wb.sheetnames) + 1)

for ws in wb.worksheets:
    hdr = 1 if ws.title == "README" else 3

    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(
                name="Times New Roman",
                size=10,
            )

            if isinstance(cell.value, float):
                cell.number_format = (
                    "0.000E+00"
                    if (
                        cell.value != 0
                        and (
                            abs(cell.value) < 1e-3
                            or abs(cell.value) >= 1e5
                        )
                    )
                    else "0.0000"
                )

    if hdr == 3:
        ws["A1"].font = Font(
            name="Times New Roman",
            size=11,
            bold=True,
        )

    for cell in ws[hdr]:
        cell.font = Font(
            name="Times New Roman",
            size=10,
            bold=True,
        )
        cell.fill = PatternFill(
            "solid",
            fgColor="DDDDDD",
        )
        cell.alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

    ws.freeze_panes = ws.cell(
        row=hdr + 1,
        column=1,
    )

    for i, col in enumerate(
        ws.iter_cols(
            min_row=hdr,
            max_row=min(ws.max_row, hdr + 60),
        ),
        start=1,
    ):
        length = max(
            len(str(cell.value))
            if cell.value is not None
            else 0
            for cell in col
        )
        ws.column_dimensions[
            get_column_letter(i)
        ].width = min(
            max(10, length + 2),
            50,
        )

wb.save(path)

print(
    path,
    [(r["sheet"], r["rows"]) for r in readme],
)