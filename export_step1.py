"""Write every step-1 table in full to the supplementary workbook."""
import json, pickle, numpy as np, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from biabak import config as C
OUT = C.OUT_DIR / "step1"
T = pickle.load(open(OUT / "step1_tables.pkl", "rb"))
order = [
 ("S1_populations", "Analytical populations and their use"),
 ("S2_extraction_check", "Re-extraction of covariates from the processed rasters compared with the extraction file of 5 August 2026"),
 ("S3_borehole_covariates", "Borehole records, UTM 32N coordinates and all covariates (224 boreholes)"),
 ("S4_extraction_flags", "Boreholes whose covariate value required the documented fallback or remained missing"),
 ("S5_descriptive_statistics", "Descriptive statistics of hydraulic variables"),
 ("S6_digit_preference", "Digit preference in recorded depth and discharge"),
 ("S7_nearest_neighbour", "Nearest recorded borehole for each borehole"),
 ("S8_close_pairs_under_2km", "All borehole pairs separated by less than 2 km"),
 ("S9_experimental_variograms_omni", "Omnidirectional experimental variograms, 2 km lags, with 999-permutation envelopes"),
 ("S10_structure_tests", "Global deviation tests and one-sided short-lag continuity tests against random relabelling"),
 ("S11_model_fits", "Variogram model fits for every target, estimator, weighting and model family"),
 ("S12_loo_kriging", "Leave-one-out ordinary kriging by model family (variogram fitted on all data; interpolation diagnostic, not a transfer test)"),
 ("S13_directional_variograms", "Directional experimental variograms (tolerance 22.5 degrees) with 499-permutation envelopes"),
 ("S14_directional_fits", "Directional model fits and global tests with Holm adjustment over 12 tests"),
 ("S15_short_range_structure", "Semivariance of sub-kilometre pairs against random relabelling"),
 ("S16_cluster_decomposition", "Variograms computed separately for pairs inside, across and outside the dense cluster"),
 ("S17_cluster_variance_comparison", "Variance of each target inside and outside the dense cluster"),
 ("S17b_structure_tests_by_subset", "Continuity tests restricted to boreholes inside or outside the dense cluster"),
 ("S17c_variograms_by_subset", "Experimental variograms and envelopes inside and outside the dense cluster"),
 ("S18_proportional_effect", "Local mean against local variance in 10 km windows"),
 ("S19_sensitivity_lag_estimator_weighting", "Fitted parameters across lag spacing, estimator and weighting"),
 ("S20_sensitivity_declustering", "Fitted parameters after cell declustering (5 to 30 km cells, 25 origins)"),
 ("S21_sensitivity_possible_duplicates", "Fitted parameters after keeping one record per pair closer than 20 m"),
]
path = OUT / "Supplementary_Tables_Step1_Variography.xlsx"
readme = []
with pd.ExcelWriter(path, engine="openpyxl") as xw:
    for n, (k, title) in enumerate(order, start=1):
        sid = f"Table S{n}"
        df = T[k].copy()
        for c in df.columns:
            if df[c].dtype == bool: df[c] = df[c].map({True: "yes", False: "no"})
        sheet = f"S{n}"
        pd.DataFrame([[f"{sid}. {title}"]]).to_excel(xw, sheet_name=sheet, index=False, header=False, startrow=0)
        df.to_excel(xw, sheet_name=sheet, index=False, startrow=2)
        readme.append(dict(table=sid, sheet=sheet, title=title, rows=len(df), columns=len(df.columns), source_key=k))
    settings = pd.DataFrame([(k, str(getattr(C, k))) for k in dir(C) if k.isupper() and not k.endswith(("_DIR", "_XLSX", "_SHP"))],
                            columns=["setting", "value"])
    settings.to_excel(xw, sheet_name="Settings", index=False)
    pd.DataFrame(json.loads((OUT / "step1_key_numbers.json").read_text()).items(), columns=["quantity", "value"]).to_excel(xw, sheet_name="Key_numbers", index=False)
    pd.DataFrame(readme).to_excel(xw, sheet_name="README", index=False)
wb = load_workbook(path)
wb.move_sheet("README", offset=-len(wb.sheetnames) + 1)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name="Times New Roman", size=10)
            if isinstance(c.value, float):
                c.number_format = "0.000E+00" if (c.value != 0 and (abs(c.value) < 1e-3 or abs(c.value) >= 1e5)) else "0.0000"
    hdr = 3 if ws.title.startswith("S") and ws.title != "Settings" else 1
    if hdr == 3:
        ws["A1"].font = Font(name="Times New Roman", size=11, bold=True)
    for c in ws[hdr]:
        c.font = Font(name="Times New Roman", size=10, bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    for i, col in enumerate(ws.iter_cols(min_row=hdr, max_row=min(ws.max_row, hdr + 60)), start=1):
        L = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, L + 2), 45)
wb.save(path)
print(path, len(order), "tables")
