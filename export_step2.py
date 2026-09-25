"""Write every step-2 table in full to the supplementary workbook."""
import pickle, numpy as np, pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from biabak import config as C, models as M
OUT = C.OUT_DIR / "step2"
T = pickle.load(open(OUT / "step2_tables.pkl", "rb"))
hp = pd.DataFrame(list(M.HYPERPARAMETERS.items()), columns=["model or step", "setting"])
designs = pd.DataFrame([
 ("random_single", "one stratified random 70/30 split", "borehole"),
 ("random_repeated_100", "100 stratified random 70/30 splits with predefined seeds", "split (distribution reported)"),
 ("kmeans_8", "8 blocks from k-means on UTM coordinates, leave one block out", "held-out block"),
 ("grid_10km", "square blocks of 10 km, leave one block out", "held-out block"),
 ("grid_20km", "square blocks of 20 km, leave one block out", "held-out block"),
 ("grid_35km", "square blocks of 35 km, leave one block out", "held-out block"),
 ("grid_50km", "square blocks of 50 km, leave one block out", "held-out block"),
 ("nndm_loo", "nearest-neighbour distance matching leave-one-out against the 1 km cells of the Centre Region (Mila et al., 2022); at least 50% of boreholes kept for training; primary estimand", "borehole"),
], columns=["design", "definition", "bootstrap resampling unit"])
seed_path = OUT / "mlp_seed_check.csv"
order = [
 ("hyper", hp, "Prespecified hyperparameters and preprocessing"),
 ("designs", designs, "Validation designs"),
 ("generalisation_distance", T["generalisation_distance"], "Distance from each held-out borehole to its nearest training borehole, by design, against the prediction domain"),
 ("fold_metadata", T["fold_metadata"], "Every fold of every design: sizes, events, distances and the variogram refitted on the training fold"),
 ("summary_reg", T["summary_by_design_model"][T["summary_by_design_model"].task == "regression"].dropna(axis=1, how="all"), "Regression of log10 Sc: skill by design and model with 95% bootstrap intervals"),
 ("summary_clf", T["summary_by_design_model"][T["summary_by_design_model"].task == "classification"].dropna(axis=1, how="all"), "Classification of unsuccessful boreholes: skill by design and model with 95% bootstrap intervals"),
 ("primary_tests", T["primary_nndm_tests"], "Primary estimand (NNDM): one-sided bootstrap tests with Holm adjustment across models"),
 ("per_split", T["per_split_random_repeated"], "Every one of the 100 random splits, every model"),
 ("per_fold", T["per_fold_block_designs"], "Every held-out block of every block design, every model"),
 ("skill_distance", T["skill_by_distance"], "Skill by distance to the nearest training borehole (pooled random, block and NNDM predictions; borehole bootstrap)"),
]
if seed_path.exists():
    order.append(("mlp_seeds", pd.read_csv(seed_path), "Sensitivity of the MLP classifier under NNDM to its random seed"))
path = OUT / "Supplementary_Tables_Step2_Validation.xlsx"
readme = []
with pd.ExcelWriter(path, engine="openpyxl") as xw:
    for n, (k, df, title) in enumerate(order, start=24):   # numbering continues after step 1
        df = df.copy()
        for c in df.columns:
            if df[c].dtype == bool: df[c] = df[c].map({True: "yes", False: "no"})
            if str(df[c].dtype) == "category": df[c] = df[c].astype(str)
        sheet = f"S{n}"
        pd.DataFrame([[f"Table S{n}. {title}"]]).to_excel(xw, sheet_name=sheet, index=False, header=False)
        df.to_excel(xw, sheet_name=sheet, index=False, startrow=2)
        readme.append(dict(table=f"Table S{n}", sheet=sheet, title=title, rows=len(df), columns=len(df.columns)))
    pd.DataFrame(readme).to_excel(xw, sheet_name="README", index=False)
wb = load_workbook(path); wb.move_sheet("README", offset=-len(wb.sheetnames) + 1)
for ws in wb.worksheets:
    hdr = 1 if ws.title == "README" else 3
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name="Times New Roman", size=10)
            if isinstance(c.value, float):
                c.number_format = "0.000E+00" if (c.value != 0 and (abs(c.value) < 1e-3 or abs(c.value) >= 1e5)) else "0.0000"
    if hdr == 3: ws["A1"].font = Font(name="Times New Roman", size=11, bold=True)
    for c in ws[hdr]:
        c.font = Font(name="Times New Roman", size=10, bold=True); c.fill = PatternFill("solid", fgColor="DDDDDD")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    for i, col in enumerate(ws.iter_cols(min_row=hdr, max_row=min(ws.max_row, hdr + 60)), start=1):
        L = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, L + 2), 50)
wb.save(path); print(path, [(r["sheet"], r["rows"]) for r in readme])
